from pathlib import Path

import pandas as pd
from pypdf import PdfReader
from docx import Document as DocxDocument
from pptx import Presentation


SUPPORTED_EXTENSIONS = {
    ".pdf", ".docx", ".txt", ".csv", ".xlsx", ".pptx"
}


def make_document(text, file_path, page_number=None):
    """Create a consistent document record with source metadata."""
    return {
        "text": text,
        "source": file_path.name,
        "file_type": file_path.suffix.lower().lstrip("."),
        "page_number": page_number,
    }


def load_documents(upload_dir="uploads"):
    """Load supported documents from the specified directory."""
    documents = []
    upload_path = Path(upload_dir)

    if not upload_path.exists():
        return documents

    for file_path in sorted(upload_path.iterdir()):
        if not file_path.is_file():
            continue

        extension = file_path.suffix.lower()

        if extension not in SUPPORTED_EXTENSIONS:
            continue

        try:
            # Extract every PDF page separately.
            if extension == ".pdf":
                reader = PdfReader(str(file_path))

                for page_number, page in enumerate(reader.pages, start=1):
                    text = page.extract_text() or ""

                    if text.strip():
                        documents.append(
                            make_document(
                                text=text,
                                file_path=file_path,
                                page_number=page_number,
                            )
                        )

            elif extension == ".docx":
                doc = DocxDocument(str(file_path))

                parts = [
                    paragraph.text
                    for paragraph in doc.paragraphs
                    if paragraph.text.strip()
                ]

                # Include table contents where present.
                for table in doc.tables:
                    for row in table.rows:
                        parts.append(
                            " | ".join(cell.text for cell in row.cells)
                        )

                text = "\n".join(parts)

                if text.strip():
                    documents.append(make_document(text, file_path))

            elif extension == ".txt":
                text = file_path.read_text(
                    encoding="utf-8",
                    errors="replace",
                )

                if text.strip():
                    documents.append(make_document(text, file_path))

            elif extension == ".csv":
                dataframe = pd.read_csv(
                    file_path,
                    dtype=str,
                    encoding="utf-8",
                ).fillna("")

                text = dataframe.to_string(index=False)

                if text.strip():
                    documents.append(make_document(text, file_path))

            elif extension == ".xlsx":
                workbook = pd.ExcelFile(file_path)
                sheet_texts = []

                for sheet_name in workbook.sheet_names:
                    dataframe = pd.read_excel(
                        workbook,
                        sheet_name=sheet_name,
                        dtype=str,
                    ).fillna("")

                    sheet_texts.append(
                        f"Sheet: {sheet_name}\n"
                        f"{dataframe.to_string(index=False)}"
                    )

                workbook.close()
                text = "\n\n".join(sheet_texts)

                if text.strip():
                    documents.append(make_document(text, file_path))

            elif extension == ".pptx":
                presentation = Presentation(str(file_path))
                slide_texts = []

                for slide_number, slide in enumerate(
                    presentation.slides, start=1
                ):
                    parts = [
                        shape.text
                        for shape in slide.shapes
                        if hasattr(shape, "text") and shape.text.strip()
                    ]

                    if parts:
                        slide_texts.append(
                            f"Slide {slide_number}\n" + "\n".join(parts)
                        )

                text = "\n\n".join(slide_texts)

                if text.strip():
                    documents.append(make_document(text, file_path))

        except Exception as error:
            # Continue loading other files if one file cannot be read.
            print(f"Could not read {file_path.name}: {error}")

    return documents